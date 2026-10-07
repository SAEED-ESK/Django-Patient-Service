from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from django.utils import timezone


def export_patients_to_excel(queryset):

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "بیماران"

    headers = [
        "شناسه",
        "سن",
        "جنس",
        "تشخیص روان پزشکی",

        "دیابت",
        "پرفشاری خون",
        "دیس‌لیپیدمی",
        "چاقی",
        "بیماری قلبی عروقی",

        "تاریخ شروع درمان",
        "آخرین تغییر دارو یا دوز",

        "پایش وزن / BMI",
        "پایش دور کمر",
        "پایش فشار خون / نبض",
        "پایش گلوکز / HbA1c",
        "پایش پروفایل چربی",
        "پایش پرولاکتین",
        "ارزیابی اختلالات حرکتی",
        "ECG",

        "انطباق کلی",

        "داروها",

        "توضیحات",

        "تاریخ ثبت",
        "آخرین بروزرسانی",
    ]

    worksheet.append(headers)

    # استایل عنوان ستون‌ها
    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # اطلاعات بیماران
    for patient in queryset:

        medications = []

        for medication in patient.medications.all():
            medications.append(
                f"{medication.get_drug_display()} - "
                f"{medication.dose} - "
                f"{medication.get_form_display()}"
            )

        worksheet.append([
            patient.id,
            patient.age,
            patient.get_gender_display(),
            patient.psychiatric_diagnosis,

            "دارد" if patient.diabetes else "ندارد",
            "دارد" if patient.hypertension else "ندارد",
            "دارد" if patient.dyslipidemia else "ندارد",
            "دارد" if patient.obesity else "ندارد",
            "دارد" if patient.cardiovascular_disease else "ندارد",

            patient.treatment_start_date,
            patient.last_medication_change,

            patient.get_weight_status(),
            patient.get_waist_status(),
            patient.get_blood_pressure_status(),
            patient.get_glucose_status(),
            patient.get_lipid_status(),
            patient.get_prolactin_status(),
            patient.get_movement_status(),
            patient.get_ecg_status(),

            patient.get_overall_compliance(),

            "\n".join(medications) if medications else "بدون دارو",

            patient.notes,

            timezone.make_naive(patient.created_at),
            timezone.make_naive(patient.updated_at),
        ])

    # تراز وسط برای تمام سلول‌ها
    for row in worksheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # فریز کردن ردیف عنوان
    worksheet.freeze_panes = "A2"

    # فعال کردن فیلتر Excel
    worksheet.auto_filter.ref = worksheet.dimensions

    # تنظیم عرض ستون‌ها
    for column in worksheet.columns:

        max_length = 0
        column_letter = column[0].column_letter

        for cell in column:
            if cell.value is not None:
                max_length = max(
                    max_length,
                    len(str(cell.value))
                )

        worksheet.column_dimensions[
            column_letter
        ].width = min(max_length + 3, 35)

    # ارتفاع ردیف‌ها
    for row in worksheet.iter_rows():
        worksheet.row_dimensions[
            row[0].row
        ].height = 30

    # ساخت فایل در حافظه
    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        'attachment; filename="patients.xlsx"'
    )

    return response